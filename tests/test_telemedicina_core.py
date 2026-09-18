from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from telemedicina_verticale.telemedicina_core import (
    generate_trace,
    validate_source_workbook,
)


CONFIG = {
    "source_sheet": "3523 AL 31_08_2026_new",
    "target_sheet": "Tracciato Prestazioni",
    "source_columns": {
        "patient_id": ["ID paziente"],
        "region": ["Regione Residenza String", "Regione di residenza"],
        "asl": ["ASL di residenza"],
        "age": ["fascia d'età"],
        "gender": ["Sesso"],
        "provider": ["Erogante", "Erogatore"],
    },
    "lookup_sheets": {
        "asl": "ASL",
        "age": "Fascia d'età",
        "gender": "Genere",
        "structure": "Codice-descrizione Struttura",
    },
    "target_columns": {
        "age": "E",
        "assistance_asl": "F",
        "gender": "G",
        "quantity": "H",
        "provider_asl": "P",
        "structure": "S",
    },
    "supplemental_asl": {},
    "structure_translation": {
        "CASTROVILLARI": {
            "provider_asl_code": "180201",
            "structure_code": "180006",
        }
    },
}


def create_source(path: Path, *, include_gender: bool = True) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Foglio dati rinominato"
    headers = [
        "ID paziente",
        "Regione di residenza",
        "ASL di residenza",
        "fascia d'età",
    ]
    if include_gender:
        headers.append("Sesso")
    headers.extend(["Erogatore", "Colonna extra non usata"])
    sheet.append(["Titolo del report"])
    sheet.append(headers)
    for patient_id in ("P1", "P2"):
        values = [patient_id, 180, 201, "18-34"]
        if include_gender:
            values.append("F")
        values.extend(["CASTROVILLARI", "da ignorare"])
        sheet.append(values)
    workbook.save(path)


def add_lookup(workbook: Workbook, title: str, rows: list[tuple[str, str, str]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append([None, "Code", "Display", "Concatenazione"])
    for code, display, concatenated in rows:
        sheet.append([None, code, display, concatenated])


def create_template(path: Path) -> None:
    workbook = Workbook()
    target = workbook.active
    target.title = "Tracciato Prestazioni"
    target.append([f"Colonna {index}" for index in range(1, 23)])
    target.append([None] * 22)
    add_lookup(workbook, "ASL", [("180201", "ASP COSENZA", "180201, ASP COSENZA")])
    add_lookup(workbook, "Fascia d'età", [("A1", "18-34", "A1, 18-34")])
    add_lookup(workbook, "Genere", [("2", "F", "2, F")])
    add_lookup(
        workbook,
        "Codice-descrizione Struttura",
        [("180006", "CASTROVILLARI", "180006, CASTROVILLARI")],
    )
    workbook.save(path)


class TelemedicinaCoreTests(unittest.TestCase):
    def test_reports_only_missing_used_columns(self) -> None:
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            create_source(source, include_gender=False)

            result = validate_source_workbook(source, CONFIG)

            self.assertFalse(result.valid)
            self.assertEqual(("gender",), result.missing_fields)
            message = result.message(CONFIG["source_columns"])
            self.assertIn("'Sesso'", message)
            self.assertNotIn("Colonna extra non usata", message)

    def test_accepts_aliases_extra_columns_and_renamed_sheet(self) -> None:
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            create_source(source)

            result = validate_source_workbook(source, CONFIG)

            self.assertTrue(result.valid)
            self.assertEqual("Foglio dati rinominato", result.sheet_name)
            self.assertEqual(2, result.header_row)

    def test_output_total_is_derived_from_source(self) -> None:
        with TemporaryDirectory() as temp_dir:
            base = Path(temp_dir)
            source = base / "source.xlsx"
            template = base / "template.xlsx"
            output = base / "output.xlsx"
            create_source(source)
            create_template(template)

            summary = generate_trace(
                source,
                output,
                template_path=template,
                cfg=CONFIG,
            )

            self.assertEqual(2, summary.total_cases)
            self.assertEqual(2, summary.output_quantity_total)
            self.assertEqual(1, summary.aggregated_rows)
            workbook = load_workbook(output, data_only=True)
            try:
                sheet = workbook["Tracciato Prestazioni"]
                self.assertEqual(2, sheet["H2"].value)
                self.assertEqual("180201, ASP COSENZA", sheet["F2"].value)
                self.assertEqual("180201, ASP COSENZA", sheet["P2"].value)
                self.assertEqual("180006, CASTROVILLARI", sheet["S2"].value)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
